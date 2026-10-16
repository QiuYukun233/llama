"""Run with:  python -m pytest archdiag/tests -q   (or  python archdiag/tests/test_archdiag.py)"""
import os
import random
import sys

sys.path.insert(0, os.path.join(os.path.dirname(__file__), "..", ".."))

import archdiag
from archdiag.parser import parse
from archdiag.layout import layout

EX = os.path.join(os.path.dirname(__file__), "..", "examples")


# ----------------------------------------------------------------- parser
def test_parser_tolerant_forms():
    g = parse("""```mermaid
    graph TD
      用户[用户 浏览器] -> CDN
      CDN -->|HTTPS| GW(网关)
      GW --> A & B
      A -- 读写 --> DB[(MySQL)]
      B -.-> C{{Redis}}
      A -> MQ: 发消息
      MQ ==> W
      W <--> B
      X → Y
      api-gateway --> Z
      subgraph 后端
        A; B
      end
    ```""")
    assert not g.errors, [str(i) for i in g.errors]
    assert g.nodes["用户"].label == "用户 浏览器"
    assert g.nodes["GW"].shape == "round" and g.nodes["DB"].shape == "cylinder"
    assert g.nodes["C"].shape == "hexagon"
    e = {(e.src, e.dst): e for e in g.edges}
    assert e[("CDN", "GW")].label == "HTTPS"
    assert e[("A", "DB")].label == "读写"
    assert e[("A", "MQ")].label == "发消息"
    assert e[("B", "C")].style == "dashed" and e[("MQ", "W")].style == "thick"
    assert e[("W", "B")].tail and e[("W", "B")].head
    assert ("X", "Y") in e and ("api-gateway", "Z") in e
    assert g.nodes["A"].group == "后端" and g.nodes["B"].group == "后端"


def test_parser_errors_are_helpful():
    g = parse("graph TD\n A[未闭合 --> B\n C --> \n end")
    msgs = "\n".join(str(i) for i in g.errors)
    assert "闭合" in msgs and "end" in msgs
    assert any(i.line_no == 2 for i in g.errors)


def test_json_input():
    g = parse('{"nodes":[{"id":"a","label":"A","group":"g"},"b"],"edges":[{"from":"a","to":"b","label":"x"},"b -> c"]}')
    assert not g.errors
    assert g.nodes["a"].group == "g" and ("b", "c") in {(e.src, e.dst) for e in g.edges}


# ----------------------------------------------------------------- layout invariants
def _rects_overlap(a, b, pad=0.5):
    return not (a[0] + a[2] <= b[0] + pad or b[0] + b[2] <= a[0] + pad
                or a[1] + a[3] <= b[1] + pad or b[1] + b[3] <= a[1] + pad)


def _seg_hits_rect(p, q, r, shrink=1.0):
    x1, y1, w, h = r[0] + shrink, r[1] + shrink, r[2] - 2 * shrink, r[3] - 2 * shrink
    if w <= 0 or h <= 0:
        return False
    (ax, ay), (bx, by) = p, q
    lo_x, hi_x = sorted((ax, bx))
    lo_y, hi_y = sorted((ay, by))
    return lo_x < x1 + w and hi_x > x1 and lo_y < y1 + h and hi_y > y1


def check_invariants(g):
    nodes = list(g.nodes.values())
    rects = {n.id: (n.x, n.y, n.w, n.h) for n in nodes}
    # 1. nodes never overlap
    for i, a in enumerate(nodes):
        for b in nodes[i + 1:]:
            assert not _rects_overlap(rects[a.id], rects[b.id]), f"nodes overlap: {a.id} / {b.id}"
    # 2. group boxes contain their members and no foreign node
    for gid, grp in g.groups.items():
        if not grp.w:
            continue
        box = (grp.x, grp.y, grp.w, grp.h)
        members = {n.id for n in g.members(gid)}
        for n in nodes:
            inside = (n.x >= grp.x - 0.01 and n.y >= grp.y - 0.01
                      and n.x + n.w <= grp.x + grp.w + 0.01 and n.y + n.h <= grp.y + grp.h + 0.01)
            touches = _rects_overlap(box, rects[n.id])
            if n.id in members:
                assert inside, f"{n.id} should be inside group {gid}"
            else:
                assert not touches, f"foreign node {n.id} touches group {gid}"
    # 3. edges: orthogonal, endpoints on the node borders, no segment through a node interior
    for e in g.edges:
        pts = e.points
        assert len(pts) >= 2, f"edge {e.src}->{e.dst} has no route"
        for (ax, ay), (bx, by) in zip(pts, pts[1:]):
            assert abs(ax - bx) < 1e-6 or abs(ay - by) < 1e-6, f"non-orthogonal segment on {e.src}->{e.dst}"
        for nid, pt in ((e.src, pts[0]), (e.dst, pts[-1])):
            x, y, w, h = rects[nid]
            on_border = (abs(pt[0] - x) < 1e-6 or abs(pt[0] - x - w) < 1e-6) and y - 1e-6 <= pt[1] <= y + h + 1e-6 \
                or (abs(pt[1] - y) < 1e-6 or abs(pt[1] - y - h) < 1e-6) and x - 1e-6 <= pt[0] <= x + w + 1e-6
            assert on_border, f"edge {e.src}->{e.dst} endpoint {pt} not on border of {nid} {rects[nid]}"
        for p, q in zip(pts, pts[1:]):
            for n in nodes:
                if n.id in (e.src, e.dst):
                    continue
                assert not _seg_hits_rect(p, q, rects[n.id]), f"edge {e.src}->{e.dst} crosses node {n.id}"
    assert g.width > 0 and g.height > 0


def test_examples_respect_invariants():
    for fn in sorted(f for f in os.listdir(EX) if f.endswith((".mmd", ".json"))):
        text = open(os.path.join(EX, fn), encoding="utf-8").read()
        for d in (None, "TB", "LR", "BT", "RL"):
            g = archdiag.build(text, direction=d)
            check_invariants(g)


def random_graph(seed):
    rnd = random.Random(seed)
    n = rnd.randint(2, 18)
    ids = [f"n{i}" for i in range(n)]
    lines = ["graph TD"]
    ngroups = rnd.randint(0, 4)
    groups = [f"g{i}" for i in range(ngroups)]
    assigned = {}
    for nid in ids:
        if groups and rnd.random() < 0.6:
            assigned[nid] = rnd.choice(groups)
    parent = {}
    for gi in groups:
        if rnd.random() < 0.3:
            cand = [x for x in groups if x != gi]
            if cand:
                parent[gi] = rnd.choice(cand)
    # avoid parent cycles: only allow parent with smaller index
    parent = {k: v for k, v in parent.items() if int(v[1:]) < int(k[1:])}

    def emit_group(gi, indent):
        lines.append(" " * indent + f"subgraph {gi}")
        for nid, gg in assigned.items():
            if gg == gi:
                lines.append(" " * (indent + 2) + f"{nid}[节点 {nid}]")
        for child, par in parent.items():
            if par == gi:
                emit_group(child, indent + 2)
        lines.append(" " * indent + "end")

    for gi in groups:
        if gi not in parent:
            emit_group(gi, 2)
    shapes = ["", "(x)", "([x])", "[(x)]", "{x}", "{{x}}", "((x))"]
    for nid in ids:
        if nid not in assigned:
            lines.append(f"  {nid}{rnd.choice(shapes).replace('x', nid)}")
    m = rnd.randint(1, min(30, n * 2))
    arrows = ["-->", "-.->", "==>", "<-->", "---", "-->|标签|"]
    for _ in range(m):
        a, b = rnd.sample(ids, 2)
        lines.append(f"  {a} {rnd.choice(arrows)} {b}")
    return "\n".join(lines)


def test_random_graphs_respect_invariants():
    for seed in range(150):
        text = random_graph(seed)
        g = parse(text)
        assert not g.errors, (seed, [str(i) for i in g.errors], text)
        for d in ("TB", "LR"):
            g = archdiag.build(text, direction=d)
            try:
                check_invariants(g)
            except AssertionError as ex:
                raise AssertionError(f"seed {seed} dir {d}: {ex}\n{text}") from ex


def test_outputs(tmp_path=None):
    import tempfile
    tmp = tmp_path or tempfile.mkdtemp()
    text = open(os.path.join(EX, "web_arch.mmd"), encoding="utf-8").read()
    svg = archdiag.to_svg(text)
    assert svg.startswith("<svg") and "MySQL" in svg
    out = os.path.join(str(tmp), "a.pptx")
    archdiag.render(text, out, direction="auto")
    from pptx import Presentation
    prs = Presentation(out)
    names = [s.name for s in prs.slides[0].shapes]
    assert any(x.startswith("node:") for x in names) and any(x.startswith("edge:") for x in names)
    assert any(x.startswith("group:") for x in names)


if __name__ == "__main__":
    for name, fn in list(globals().items()):
        if name.startswith("test_") and callable(fn):
            fn()
            print("ok", name)
