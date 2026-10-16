"""PPTX renderer built on python-pptx.

Nodes are real autoshapes (editable text), groups are rectangles, arrows are
connectors (2 points) or freeform polylines (3+ points) with proper arrow
heads — no guessing of flipH/flipV/rot, which is where LLM-written
python-pptx code usually goes wrong.
"""
from __future__ import annotations

from typing import List, Optional, Sequence, Tuple

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.dml import MSO_LINE
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, MSO_AUTO_SIZE, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Emu, Inches, Pt

from .model import Graph, Node
from .theme import DEFAULT_THEME, Theme

EMU_PER_PX = 9525          # 96 dpi

SHAPE_MAP = {
    "rect": MSO_SHAPE.RECTANGLE,
    "round": MSO_SHAPE.ROUNDED_RECTANGLE,
    "stadium": MSO_SHAPE.ROUNDED_RECTANGLE,
    "cylinder": MSO_SHAPE.CAN,
    "diamond": MSO_SHAPE.DIAMOND,
    "hexagon": MSO_SHAPE.HEXAGON,
    "circle": MSO_SHAPE.OVAL,
    "subroutine": MSO_SHAPE.FLOWCHART_PREDEFINED_PROCESS,
    "parallelogram": MSO_SHAPE.PARALLELOGRAM,
    "flag": MSO_SHAPE.PENTAGON,
}


def _rgb(hex_color: str) -> RGBColor:
    return RGBColor.from_string(hex_color.lstrip("#"))


def _set_text(shape, lines: Sequence[str], size_pt: float, color: str, bold=False,
              align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE, margin_px: float = 4, font_name: Optional[str] = None):
    tf = shape.text_frame
    tf.word_wrap = True
    tf.auto_size = MSO_AUTO_SIZE.NONE
    tf.vertical_anchor = anchor
    m = Emu(int(margin_px * EMU_PER_PX))
    tf.margin_left = tf.margin_right = m
    tf.margin_top = tf.margin_bottom = Emu(int(m * 0.6))
    for i, line in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        run = p.add_run()
        run.text = line
        run.font.size = Pt(max(size_pt, 6))
        run.font.bold = bold
        run.font.color.rgb = _rgb(color)
        if font_name:
            run.font.name = font_name
            rpr = run._r.get_or_add_rPr()
            ea = rpr.find(qn("a:ea"))
            if ea is None:
                ea = rpr.makeelement(qn("a:ea"), {})
                rpr.append(ea)
            ea.set("typeface", font_name)


def _style_line(line, color: str, width_pt: float, dashed: bool, head: bool, tail: bool) -> None:
    line.color.rgb = _rgb(color)
    line.width = Pt(width_pt)
    if dashed:
        line.dash_style = MSO_LINE.DASH
    ln = line._get_or_add_ln()
    # schema order inside <a:ln>: fill, prstDash, join, headEnd, tailEnd
    for tag in ("a:headEnd", "a:tailEnd"):
        el = ln.find(qn(tag))
        if el is not None:
            ln.remove(el)
    if tail:
        ln.append(ln.makeelement(qn("a:headEnd"), {"type": "triangle", "w": "med", "len": "med"}))
    if head:
        ln.append(ln.makeelement(qn("a:tailEnd"), {"type": "triangle", "w": "med", "len": "med"}))


def add_diagram(slide, g: Graph, left, top, width, height, theme: Theme = DEFAULT_THEME,
                font_size: float = 13, line_h: float = 17, max_scale: float = 1.25,
                font_name: Optional[str] = None) -> List:
    """Draw a laid-out graph into the box (EMU values) on `slide`. Returns shapes."""
    th = theme
    left, top, width, height = int(left), int(top), int(width), int(height)
    if g.width <= 0 or g.height <= 0:
        return []
    scale = min(width / (g.width * EMU_PER_PX), height / (g.height * EMU_PER_PX), max_scale)
    k = EMU_PER_PX * scale
    ox = left + (width - g.width * k) / 2
    oy = top + (height - g.height * k) / 2
    pt = 0.75 * scale                      # px → pt at 96 dpi

    def X(x: float) -> int:
        return int(ox + x * k)

    def Y(y: float) -> int:
        return int(oy + y * k)

    def S(v: float) -> int:
        return int(v * k)

    shapes = []
    if g.title:
        tb = slide.shapes.add_textbox(X(0), Y(6), S(g.width), S(font_size + 16))
        _set_text(tb, [g.title], (font_size + 5) * pt, th.text, bold=True, font_name=font_name)
        shapes.append(tb)
    for grp in sorted(g.groups.values(), key=lambda x: x.depth):
        if not grp.w:
            continue
        box = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, X(grp.x), Y(grp.y), S(grp.w), S(grp.h))
        box.adjustments[0] = min(0.08, 6 / min(grp.w, grp.h))
        box.shadow.inherit = False
        box.fill.solid()
        box.fill.fore_color.rgb = _rgb(th.group_fills[(grp.depth - 1) % len(th.group_fills)])
        box.line.color.rgb = _rgb(th.group_stroke)
        box.line.width = Pt(0.9 * scale + 0.1)
        box.line.dash_style = MSO_LINE.DASH
        _set_text(box, [grp.label], (font_size - 1) * pt, th.group_text, bold=True,
                  align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, margin_px=8, font_name=font_name)
        box.name = f"group:{grp.id}"
        shapes.append(box)
    for n in g.nodes.values():
        shp = slide.shapes.add_shape(SHAPE_MAP.get(n.shape, MSO_SHAPE.RECTANGLE), X(n.x), Y(n.y), S(n.w), S(n.h))
        if n.shape == "round":
            shp.adjustments[0] = 0.18
        elif n.shape == "stadium":
            shp.adjustments[0] = 0.5
        elif n.shape == "hexagon":
            shp.adjustments[0] = 0.18
        elif n.shape == "cylinder":
            shp.adjustments[0] = 0.18
        shp.shadow.inherit = False
        shp.fill.solid()
        shp.fill.fore_color.rgb = _rgb(th.node_fill)
        shp.line.color.rgb = _rgb(th.node_stroke)
        shp.line.width = Pt(1.1 * scale + 0.1)
        _set_text(shp, n.lines, font_size * pt, th.text, margin_px=3, font_name=font_name)
        shp.name = f"node:{n.id}"
        shapes.append(shp)
    for e in g.edges:
        pts = [(X(x), Y(y)) for x, y in e.points]
        if len(pts) < 2:
            continue
        if len(pts) == 2:
            con = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, pts[0][0], pts[0][1], pts[1][0], pts[1][1])
            line = con.line
            shp = con
        else:
            ff = slide.shapes.build_freeform(pts[0][0], pts[0][1], scale=1.0)
            ff.add_line_segments(pts[1:], close=False)
            shp = ff.convert_to_shape()
            shp.fill.background()
            shp.shadow.inherit = False
            line = shp.line
        _style_line(line, th.edge, (2.2 if e.style == "thick" else 1.3) * scale + 0.1,
                    e.style == "dashed", e.head, e.tail)
        shp.name = f"edge:{e.src}->{e.dst}"
        shapes.append(shp)
    for e in g.edges:
        if e.label and e.label_pos:
            lx, ly = e.label_pos
            lw, lh = e.label_size
            tb = slide.shapes.add_textbox(X(lx - lw / 2), Y(ly - lh / 2), S(lw), S(lh))
            tb.fill.solid()
            tb.fill.fore_color.rgb = _rgb(th.label_bg)
            _set_text(tb, e.label.split("\n"), (font_size - 2) * pt, th.edge_text, margin_px=1, font_name=font_name)
            tb.name = f"label:{e.src}->{e.dst}"
            shapes.append(tb)
    return shapes


def to_pptx(g: Graph, path: Optional[str] = None, prs: Optional[Presentation] = None,
            theme: Theme = DEFAULT_THEME, slide_size: Tuple[float, float] = (13.333, 7.5),
            margin_in: float = 0.4, font_name: Optional[str] = None, **kw) -> Presentation:
    """Create (or extend) a presentation with one slide holding the diagram."""
    if prs is None:
        prs = Presentation()
        prs.slide_width = Inches(slide_size[0])
        prs.slide_height = Inches(slide_size[1])
    layout_idx = 6 if len(prs.slide_layouts) > 6 else len(prs.slide_layouts) - 1
    slide = prs.slides.add_slide(prs.slide_layouts[layout_idx])
    m = Inches(margin_in)
    add_diagram(slide, g, m, m, prs.slide_width - 2 * m, prs.slide_height - 2 * m,
                theme=theme, font_name=font_name, **kw)
    if path:
        prs.save(path)
    return prs
